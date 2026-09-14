"""Excel export of the structured irAE analysis results.

The queue workers emit one structured report dict per patient (see
``IraeLayer3QueueWorker.patient_structured``); this module flattens those
reports into one spreadsheet row per irAE finding and writes an ``.xlsx``
workbook at the path the user chooses.  No database is involved.
"""

from __future__ import annotations

from typing import Any

# (result key, Excel header) — stable column order for the workbook.
EXCEL_COLUMNS: list[tuple[str, str]] = [
    ("patient_id", "Paziente"),
    ("organ", "Organi di origine"),
    ("irAE_type", "Tipo irAE"),
    ("ctcae_grade", "Grado CTCAE"),
    ("first_onset_date", "Prima insorgenza"),
    ("probability_immune", "Probabilità immuno-correlata"),
    ("new_onset_vs_exacerbation", "Nuova vs riacutizzazione"),
    ("alternative_causes", "Cause alternative"),
    ("confidence", "Confidenza"),
    ("key_evidence_ids", "Evidenze chiave"),
    ("notes", "Approfondimento clinico"),
    ("immunotherapy_start", "Inizio immunoterapia"),
    ("candidates_total", "N° candidati"),
    ("analyzed_at", "Data analisi"),
    ("removed", "Rimosso"),
]


def _immunotherapy_label(report: dict[str, Any]) -> str:
    anchor = report.get("anchor")
    if not isinstance(anchor, dict):
        return ""
    first_drug = anchor.get("first_drug", "")
    first_date = anchor.get("first_date", "")
    if first_drug and first_date:
        return f"{first_drug} dal {first_date}"
    return first_drug or first_date or ""


def _finding_row(
    item: dict[str, Any],
    patient_id: str,
    report: dict[str, Any],
    removed: str,
) -> dict[str, Any]:
    organs = item.get("source_organs") or []
    organ = ", ".join(organs) if organs else item.get("organ", "")
    return {
        "patient_id": patient_id,
        "organ": organ,
        "irAE_type": item.get("irAE_type", ""),
        "ctcae_grade": item.get("ctcae_grade", ""),
        "first_onset_date": item.get("first_onset_date", ""),
        "probability_immune": item.get("probability_immune", ""),
        "new_onset_vs_exacerbation": item.get(
            "new_onset_vs_exacerbation", ""
        ),
        "alternative_causes": item.get("alternative_causes", ""),
        "confidence": item.get("confidence", ""),
        "key_evidence_ids": ", ".join(
            str(eid).strip().lstrip("#")
            for eid in (item.get("key_evidence_ids") or [])
        ),
        "notes": item.get("notes", ""),
        "immunotherapy_start": _immunotherapy_label(report),
        "candidates_total": report.get("candidates_total", ""),
        "analyzed_at": report.get("analyzed_at", ""),
        "removed": removed,
    }


def irae_rows_for_excel(
    results: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[str]]:
    """Flatten the structured reports into one row per irAE finding.

    ``results`` are the queue results ``{patient_id, markdown, error,
    structured}``; ``structured`` is ``None`` for the classic method or for
    patients whose analysis failed.  Returns ``(rows, skipped)`` where each
    row carries the ``EXCEL_COLUMNS`` keys and ``skipped`` lists the patients
    with no exportable structured finding (active OR removed).
    """
    rows: list[dict[str, Any]] = []
    skipped: list[str] = []
    for result in results:
        report = result.get("structured")
        if not isinstance(report, dict):
            skipped.append(result["patient_id"])
            continue
        findings = report.get("iraes") or []
        consolidation = report.get("consolidation") or {}
        removed_findings = consolidation.get("removed") or []
        for item in findings:
            rows.append(_finding_row(item, result["patient_id"], report, ""))
        for item in removed_findings:
            reason = item.get("removed_reason") or ""
            rows.append(_finding_row(
                item, result["patient_id"], report,
                f"rimosso — {reason}" if reason else "rimosso",
            ))
        if not findings and not removed_findings:
            skipped.append(result["patient_id"])
    return rows, skipped


def write_irae_xlsx(
    results: list[dict[str, Any]], path: str
) -> tuple[int, list[str]]:
    """Write one ``.xlsx`` workbook (sheet ``irAE``) with one row per
    finding.  Returns ``(rows_written, skipped_patients)``."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    rows, skipped = irae_rows_for_excel(results)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "irAE"
    sheet.append([header for _, header in EXCEL_COLUMNS])
    for row in rows:
        sheet.append([row[key] for key, _ in EXCEL_COLUMNS])
    for index, (_, header) in enumerate(EXCEL_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(
            14, len(header) + 4
        )
    workbook.save(path)
    return len(rows), skipped
