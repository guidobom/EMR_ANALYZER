"""Rebuild persistent irAE reports from an Excel export + the project registry.

Before the per-patient ``irae_report.json`` persistence existed, a finished
queue lived only in memory: once the results window closed, the structured
reports were gone.  This tool reconstructs the report dict the inspector
needs from two sources:

* the ``.xlsx`` exported with "📊 Esporta Excel" (``write_irae_xlsx``) — the
  final (Layer 4) irAE findings with their editable fields and the cited
  ``key_evidence_ids``;
* the project registry (``emr_registry.db``) — the atomic evidence from which
  ``anchor``, the Layer 2 ``candidates`` and the compact ``evidence``
  provenance (page, bbox, passage) are recomputed deterministically.

The rebuilt report is written with the same per-patient persistence the queue
now uses (``irae_report.json`` under the workspace), so "📂 Riepilogo irAE"
reopens it and manual corrections work exactly as for a fresh analysis.

Works on any project: the target workspace and registry are derived from
``--project`` (a folder under ``--projects-root``) and can be overridden with
``--workspace`` / ``--registry``.

Usage::

    python tools/irae_rebuild_from_xlsx.py --xlsx <file.xlsx> --project RENE
    python tools/irae_rebuild_from_xlsx.py --xlsx <file.xlsx> --project MELANOMA

Known losses vs. a fresh run: the suspects (``consolidation.suspects``) and
the per-organ LLM detail (``organ_results``) are not exported to Excel, so
they cannot be recovered.
"""

from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from dataclasses import asdict
from pathlib import Path
from typing import Any

from emr_analyzer.clinical.irae_export import EXCEL_COLUMNS
from emr_analyzer.clinical.irae_layers import (
    DEFAULT_MAX_CANDIDATES,
    _compact_evidence,
    find_ici_anchor,
    scan_for_irae,
    summarize_candidates,
)
from emr_analyzer.clinical.irae_reports import save_report

DEFAULT_PROJECTS_ROOT = Path.home() / "Desktop"

# Excel header (EXCEL_COLUMNS second element) → result key, so the tool is
# robust to the exact column order the export writes.
_HEADER_TO_KEY = {header: key for key, header in EXCEL_COLUMNS}


# ---------------------------------------------------------------------------
# Registry → deterministic layers
# ---------------------------------------------------------------------------


def _parse_bbox(raw: Any) -> list | None:
    """``bbox_json`` column → the ``bbox`` list ``_compact_evidence`` wants."""
    if isinstance(raw, (list, tuple)):
        return list(raw)
    if isinstance(raw, str) and raw.strip():
        try:
            parsed = json.loads(raw)
        except ValueError:
            return None
        return list(parsed) if isinstance(parsed, (list, tuple)) else None
    return None


def load_registry_rows(db_path: Path, patient_id: str) -> list[dict[str, Any]]:
    """Atomic evidence rows shaped like ``evidence_rows_from_models`` output.

    The ``clinical_evidence`` table already carries ``source_page``,
    ``source_text`` and ``bbox_json`` at row level (provenance the inspector
    needs to open and highlight the original PDF), so no ``data_json``
    unpacking is required for those.
    """
    con = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    try:
        records = con.execute(
            "SELECT evidence_id, category, normalized_entity, observed_date, "
            "       data_json, document_id, value_text, numeric_value, unit, "
            "       source_page, source_text, bbox_json "
            "FROM clinical_evidence WHERE patient_id = ? "
            "ORDER BY observed_date, id",
            (patient_id,),
        ).fetchall()
    finally:
        con.close()

    rows: list[dict[str, Any]] = []
    for record in records:
        row = dict(record)
        row["bbox"] = _parse_bbox(row.get("bbox_json"))
        row.pop("bbox_json", None)
        raw = row.get("data_json")
        if isinstance(raw, str) and raw.strip():
            try:
                row["data_json"] = json.loads(raw)
            except ValueError:
                row["data_json"] = None
        else:
            row["data_json"] = None
        rows.append(row)
    return rows


# ---------------------------------------------------------------------------
# Excel → findings
# ---------------------------------------------------------------------------


def rows_from_excel(xlsx_path: Path) -> list[dict[str, Any]]:
    """Flatten the exported workbook into one dict per irAE row."""
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration:
        raise ValueError(f"Excel vuoto: {xlsx_path}")

    columns: dict[str, int] = {}
    for index, name in enumerate(header):
        key = _HEADER_TO_KEY.get(str(name).strip())
        if key is not None:
            columns[key] = index
    missing = {"patient_id", "irAE_type"} - set(columns)
    if missing:
        raise ValueError(
            "Colonne mancanti nell'Excel: " + ", ".join(sorted(missing))
            + " — assicurati di esportare con '📊 Esporta Excel'."
        )

    rows: list[dict[str, Any]] = []

    def cell(key: str, default: Any = "") -> Any:
        index = columns.get(key)
        if index is None or index >= len(values):
            return default
        value = values[index]
        return default if value is None else value

    for values in iterator:
        if values is None or all(v is None or v == "" for v in values):
            continue
        rows.append({
            "patient_id": str(cell("patient_id")).strip(),
            "organ": str(cell("organ")),
            "irAE_type": str(cell("irAE_type")),
            "ctcae_grade": str(cell("ctcae_grade")),
            "first_onset_date": str(cell("first_onset_date")),
            "probability_immune": str(cell("probability_immune")),
            "new_onset_vs_exacerbation": str(cell("new_onset_vs_exacerbation")),
            "alternative_causes": str(cell("alternative_causes")),
            "confidence": cell("confidence"),
            "key_evidence_ids": str(cell("key_evidence_ids")),
            "notes": str(cell("notes")),
            "analyzed_at": str(cell("analyzed_at")),
        })
    return rows


def finding_from_row(row: dict[str, Any]) -> dict[str, Any]:
    """Rebuild one Layer 4 finding dict from its exported Excel row."""
    organs = [
        part.strip()
        for part in str(row.get("organ") or "").split(",")
        if part.strip()
    ]
    evidence_ids = [
        part.strip().lstrip("#")
        for part in str(row.get("key_evidence_ids") or "").split(",")
        if part.strip()
    ]
    finding: dict[str, Any] = {"irAE_type": str(row.get("irAE_type") or "")}
    if organs:
        finding["source_organs"] = organs
        finding["organ"] = organs[0]
    for key in (
        "ctcae_grade", "first_onset_date", "probability_immune",
        "new_onset_vs_exacerbation", "alternative_causes", "notes",
    ):
        value = str(row.get(key) or "").strip()
        if value:
            finding[key] = value
    confidence = row.get("confidence")
    if isinstance(confidence, (int, float)):
        finding["confidence"] = confidence
    elif isinstance(confidence, str) and confidence.strip():
        try:
            finding["confidence"] = float(confidence)
        except ValueError:
            finding["confidence"] = confidence.strip()
    if evidence_ids:
        finding["key_evidence_ids"] = evidence_ids
    return finding


# ---------------------------------------------------------------------------
# Report reconstruction
# ---------------------------------------------------------------------------


def rebuild_report(
    rows: list[dict[str, Any]],
    findings: list[dict[str, Any]],
    *,
    analyzed_at: str = "",
    max_candidates: int = DEFAULT_MAX_CANDIDATES,
) -> dict[str, Any]:
    """Reconstruct a structured report dict from registry rows + findings.

    Layers 1-2 (``anchor``, ``candidates``) are recomputed deterministically;
    ``evidence`` keeps only the cited ids with full provenance (as the
    pipeline does).  The consolidation is marked as applied over the rebuilt
    findings; ``organ_results`` cannot be recovered.
    """
    anchor = find_ici_anchor(rows)
    candidates = scan_for_irae(rows, anchor)
    baseline_excluded = 0
    if anchor is not None:
        baseline_excluded = sum(
            1
            for candidate in scan_for_irae(rows, anchor, drop_baseline=False)
            if candidate.band == "pre_ici"
        )

    report: dict[str, Any] = {
        "anchor": (
            {
                "first_drug": anchor.first_drug,
                "first_date": anchor.first_date.isoformat(),
                "first_raw": anchor.first_raw,
                "last_drug": anchor.last_drug,
                "last_date": anchor.last_date.isoformat(),
                "last_raw": anchor.last_raw,
                "occurrences": anchor.occurrences,
            }
            if anchor is not None
            else None
        ),
        "immunotherapy_start": (
            {
                "date": anchor.first_date.isoformat(),
                "raw": anchor.first_raw,
                "drug": anchor.first_drug,
            }
            if anchor is not None
            else None
        ),
        "candidates_total": len(candidates),
        "candidates_by_organ": summarize_candidates(candidates),
        "baseline_excluded": baseline_excluded,
        "iraes": findings,
        "consolidation": {
            "iraes": findings,
            "suspects": [],
            "applied": True,
            "error": None,
            "input_count": len(findings),
            "output_count": len(findings),
        },
        "organ_results": {},
        "candidates": [asdict(candidate) for candidate in candidates[:max_candidates]],
    }

    cited_ids: list[str] = []
    for finding in findings:
        cited_ids.extend(
            str(eid).strip().lstrip("#")
            for eid in (finding.get("key_evidence_ids") or [])
        )
    evidence_by_id = {str(row.get("evidence_id")): row for row in rows}
    report["evidence"] = _compact_evidence([
        evidence_by_id[eid]
        for eid in dict.fromkeys(cited_ids)
        if eid in evidence_by_id
    ])

    if analyzed_at:
        report["analyzed_at"] = analyzed_at
    return report


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Ricostruisce i report irAE persistiti (irae_report.json) da un "
            "Excel esportato + il registro del progetto."
        )
    )
    parser.add_argument(
        "--xlsx", type=Path, required=True,
        help="Excel esportato con '📊 Esporta Excel' dalla finestra dei risultati.",
    )
    parser.add_argument(
        "--project", required=True,
        help=(
            "Cartella del progetto, es. RENE, MELANOMA o LUNG. Obbligatoria: "
            "i report vengono scritti in <projects-root>/<project>/<paziente>/."
        ),
    )
    parser.add_argument(
        "--projects-root", type=Path, default=DEFAULT_PROJECTS_ROOT,
        help="Directory che contiene le cartelle progetto (default: Desktop).",
    )
    parser.add_argument(
        "--workspace", type=Path, default=None,
        help="Dove scrivere i report (default: la cartella del progetto).",
    )
    parser.add_argument(
        "--registry", type=Path, default=None,
        help="Percorso di emr_registry.db (default: <workspace>/emr_registry.db).",
    )
    parser.add_argument(
        "--max-candidates", type=int, default=DEFAULT_MAX_CANDIDATES,
        help=f"Limite candidati nel report (default: {DEFAULT_MAX_CANDIDATES}).",
    )
    args = parser.parse_args(argv)

    workspace = args.workspace or (args.projects_root / args.project)
    registry = args.registry or (workspace / "emr_registry.db")
    if not args.xlsx.exists():
        print(f"Excel non trovato: {args.xlsx}", file=sys.stderr)
        return 2
    if not registry.exists():
        print(f"Registro non trovato: {registry}", file=sys.stderr)
        return 2

    try:
        rows = rows_from_excel(args.xlsx)
    except ValueError as exc:
        print(f"Errore nell'Excel: {exc}", file=sys.stderr)
        return 2

    by_patient: dict[str, list[dict[str, Any]]] = {}
    analyzed_at_by_patient: dict[str, str] = {}
    for row in rows:
        pid = row["patient_id"]
        if not pid:
            continue
        by_patient.setdefault(pid, []).append(row)
        if row["analyzed_at"] and pid not in analyzed_at_by_patient:
            analyzed_at_by_patient[pid] = row["analyzed_at"]

    if not by_patient:
        print("Nessuna riga con un paziente valido nell'Excel.", file=sys.stderr)
        return 2

    from emr_analyzer.config import active_workspace

    active_workspace.set_path(workspace)

    written = 0
    warnings: list[str] = []
    for pid in sorted(by_patient):
        findings = [
            finding_from_row(row)
            for row in by_patient[pid]
            if finding_from_row(row).get("irAE_type")
        ]
        if not findings:
            warnings.append(f"{pid}: nessun irAE con tipo valido, saltato.")
            continue
        evidence_rows = load_registry_rows(registry, pid)
        if not evidence_rows:
            warnings.append(
                f"{pid}: nessuna evidenza nel registro — candidati/evidenze "
                "vuote nel report (la vista evidenze non sarà disponibile)."
            )
        report = rebuild_report(
            evidence_rows,
            findings,
            analyzed_at=analyzed_at_by_patient.get(pid, ""),
            max_candidates=args.max_candidates,
        )
        save_report(pid, report)
        written += 1
        print(
            f"[ok] {pid}: {len(findings)} irAE, {len(evidence_rows)} evidenze "
            f"dal registro → {workspace / pid / 'irae_report.json'}"
        )

    for warning in warnings:
        print(f"[avviso] {warning}", file=sys.stderr)
    print(f"Fatto: {written} report scritti in {workspace}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
