"""Tests for the Excel export of structured irAE results."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path


def _finding(**overrides) -> dict:
    base = {
        "organ": "Miocardite/Cardiotossicità",
        "irAE_type": "Elevazione della troponina",
        "ctcae_grade": "G2",
        "first_onset_date": "2022-11-22",
        "probability_immune": "PROBABILE",
        "new_onset_vs_exacerbation": "nuova insorgenza",
        "alternative_causes": "nessuna",
        "confidence": 0.9,
        "key_evidence_ids": ["E-TROP"],
    }
    base.update(overrides)
    return base


def _report(findings: list[dict], *, candidates: int = 3) -> dict:
    return {
        "anchor": {
            "first_drug": "nivolumab", "first_date": "2022-09-01",
            "first_raw": "2022-09-01", "last_drug": "nivolumab",
            "last_date": "2022-09-01", "last_raw": "2022-09-01",
            "occurrences": 1,
        },
        "candidates_total": candidates,
        "iraes": findings,
        "analyzed_at": "2026-08-27T10:00:00",
    }


def _result(patient_id: str, structured) -> dict:
    return {
        "patient_id": patient_id, "label": patient_id,
        "markdown": "# analisi", "error": None, "structured": structured,
    }


class IraeRowsForExcelTest(unittest.TestCase):
    def test_one_row_per_finding_with_all_columns(self):
        from emr_analyzer.clinical.irae_export import (
            EXCEL_COLUMNS,
            irae_rows_for_excel,
        )

        results = [_result("P001", _report([
            _finding(),
            _finding(
                organ="Oculari", irAE_type="Uveite anteriore bilaterale",
                ctcae_grade="G1", first_onset_date="2024-11-18",
                probability_immune="POSSIBILE", confidence=0.7,
                key_evidence_ids=["E-UVE", "E-9"],
            ),
        ]))]
        rows, skipped = irae_rows_for_excel(results)

        self.assertEqual(skipped, [])
        self.assertEqual(len(rows), 2)
        row = rows[0]
        self.assertEqual(row["patient_id"], "P001")
        self.assertEqual(row["organ"], "Miocardite/Cardiotossicità")
        self.assertEqual(row["ctcae_grade"], "G2")
        self.assertEqual(row["first_onset_date"], "2022-11-22")
        self.assertEqual(row["probability_immune"], "PROBABILE")
        self.assertEqual(row["key_evidence_ids"], "E-TROP")
        self.assertEqual(row["immunotherapy_start"], "nivolumab dal 2022-09-01")
        self.assertEqual(row["candidates_total"], 3)
        self.assertEqual(row["analyzed_at"], "2026-08-27T10:00:00")
        self.assertEqual(rows[1]["key_evidence_ids"], "E-UVE, E-9")
        # every result key has a header column.
        self.assertEqual(len(EXCEL_COLUMNS), 14)

    def test_source_organs_and_notes_are_exported(self):
        from emr_analyzer.clinical.irae_export import irae_rows_for_excel

        finding = _finding(
            source_organs=["Miocardite/Cardiotossicità", "Dermatite"],
            notes="ripresa con prednisone 1 mg/kg",
        )
        rows, skipped = irae_rows_for_excel([_result("P001", _report([finding]))])
        self.assertEqual(skipped, [])
        self.assertEqual(rows[0]["organ"], "Miocardite/Cardiotossicità, Dermatite")
        self.assertEqual(rows[0]["notes"], "ripresa con prednisone 1 mg/kg")

    def test_patients_without_structured_findings_are_skipped(self):
        from emr_analyzer.clinical.irae_export import irae_rows_for_excel

        results = [
            _result("P001", _report([_finding()])),
            _result("P002", None),                     # protocollo classico
            _result("P003", {"iraes": [], "anchor": None}),  # nessun irAE
            _result("P004", {"iraes": None}),          # schema anomalo
        ]
        rows, skipped = irae_rows_for_excel(results)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["patient_id"], "P001")
        self.assertEqual(skipped, ["P002", "P003", "P004"])

    def test_missing_anchor_yields_empty_immunotherapy(self):
        from emr_analyzer.clinical.irae_export import irae_rows_for_excel

        report = _report([_finding()])
        report["anchor"] = None
        rows, _ = irae_rows_for_excel([_result("P001", report)])
        self.assertEqual(rows[0]["immunotherapy_start"], "")


class WriteIraeXlsxTest(unittest.TestCase):
    def test_writes_workbook_with_header_and_rows(self):
        from openpyxl import load_workbook

        from emr_analyzer.clinical.irae_export import write_irae_xlsx

        results = [_result("P001", _report([_finding(), _finding(organ="Oculari")]))]
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "out.xlsx")
            written, skipped = write_irae_xlsx(results, path)

            self.assertEqual(written, 2)
            self.assertEqual(skipped, [])
            self.assertTrue(Path(path).exists())

            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["irAE"])
            sheet = workbook["irAE"]
            self.assertEqual(sheet.max_row, 3)  # header + 2 findings
            self.assertEqual(sheet["A1"].value, "Paziente")
            self.assertEqual(sheet["B1"].value, "Organi di origine")
            self.assertEqual(sheet["C1"].value, "Tipo irAE")
            self.assertEqual(sheet["G1"].value, "Nuova vs riacutizzazione")
            self.assertEqual(sheet["K1"].value, "Approfondimento clinico")
            self.assertEqual(sheet["A2"].value, "P001")
            self.assertEqual(sheet["B2"].value, "Miocardite/Cardiotossicità")
            self.assertEqual(sheet["D2"].value, "G2")
            self.assertEqual(sheet["B3"].value, "Oculari")

    def test_skipped_patients_are_reported(self):
        from openpyxl import load_workbook

        from emr_analyzer.clinical.irae_export import write_irae_xlsx

        results = [
            _result("P001", _report([_finding()])),
            _result("P002", None),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "out.xlsx")
            written, skipped = write_irae_xlsx(results, path)
            self.assertEqual(written, 1)
            self.assertEqual(skipped, ["P002"])
            workbook = load_workbook(path)
            self.assertEqual(workbook["irAE"].max_row, 2)  # header + 1


if __name__ == "__main__":
    unittest.main()
